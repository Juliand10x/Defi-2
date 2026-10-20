import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "06_presentacion"
GRAFICAS = OUT_DIR / "graficas"

def build_presentation():
    prs = Presentation()
    # Set slide width and height to 16:9 (13.33 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    blank_layout = prs.slide_layouts[6] # Blank slide
    slide = prs.slides.add_slide(blank_layout)
    
    # Background color (Dark Navy: #0F172A)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(15, 23, 42)
    bg.line.fill.background()
    
    # Header Banner Box
    header = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.4), Inches(0.3), Inches(12.533), Inches(0.85))
    header.fill.solid()
    header.fill.fore_color.rgb = RGBColor(30, 41, 59)
    header.line.color.rgb = RGBColor(51, 65, 85)
    
    tf = header.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_top = Inches(0.1)
    
    p = tf.paragraphs[0]
    p.text = "ONE-PAGER CAPM: SÍNTESIS INTEGRADA DE LOS 10 EXPEDIENTES"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = RGBColor(96, 165, 250)
    
    p2 = tf.add_paragraph()
    p2.text = "La Beta de los Pobres (SOL vs NCI) y Acciones de Bajo Precio frente a ETF Sectoriales | Cliente: Doña Victoria | Estudiante: Julián Duarte"
    p2.font.size = Pt(11)
    p2.font.color.rgb = RGBColor(148, 163, 184)
    
    # 4 KPI Boxes
    kpis = [
        ("UNIVERSO ANALIZADO", "10 Activos", "1 Cripto + 9 Acciones en 3 Sectores", RGBColor(59, 130, 246)),
        ("MÁXIMA BETA CAPM", "1.94 (KLXE)", "Sensibilidad extrema de mercado", RGBColor(139, 92, 246)),
        ("PEOR MAX DRAWDOWN", "-93.9% (MVIS)", "Destrucción severa de capital", RGBColor(239, 68, 68)),
        ("EXCESO RETORNO > CAPM", "4 / 10 Activos", "KOPN (+38.7pp), LIDR, CTMX, TTI", RGBColor(16, 185, 129)),
    ]
    
    kpi_w = Inches(2.98)
    kpi_h = Inches(0.85)
    kpi_y = Inches(1.25)
    
    for i, (label, val, sub, color) in enumerate(kpis):
        kpi_x = Inches(0.4 + i * 3.18)
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, kpi_x, kpi_y, kpi_w, kpi_h)
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(30, 41, 59)
        box.line.color.rgb = color
        
        btf = box.text_frame
        btf.word_wrap = True
        btf.margin_left = Inches(0.12)
        btf.margin_top = Inches(0.08)
        
        bp1 = btf.paragraphs[0]
        bp1.text = label
        bp1.font.size = Pt(9)
        bp1.font.bold = True
        bp1.font.color.rgb = color
        
        bp2 = btf.add_paragraph()
        bp2.text = val
        bp2.font.size = Pt(16)
        bp2.font.bold = True
        bp2.font.color.rgb = RGBColor(248, 250, 252)
        
        bp3 = btf.add_paragraph()
        bp3.text = sub
        bp3.font.size = Pt(8.5)
        bp3.font.color.rgb = RGBColor(148, 163, 184)
        
    # Content Columns (Y = 2.22)
    # Left Column: Brecha CAPM Image (Width: 3.9 in)
    x_left = Inches(0.4)
    y_content = Inches(2.22)
    w_col = Inches(3.9)
    h_col = Inches(4.5)
    
    box_left = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_left, y_content, w_col, h_col)
    box_left.fill.solid()
    box_left.fill.fore_color.rgb = RGBColor(30, 41, 59)
    box_left.line.color.rgb = RGBColor(51, 65, 85)
    
    ltf = box_left.text_frame
    ltf.margin_left = Inches(0.12)
    ltf.margin_top = Inches(0.1)
    lp = ltf.paragraphs[0]
    lp.text = "1. BRECHA RETORNO OBSERVADO VS CAPM"
    lp.font.size = Pt(11)
    lp.font.bold = True
    lp.font.color.rgb = RGBColor(96, 165, 250)
    
    img_brecha = GRAFICAS / "07_brecha_retorno_capm.png"
    if img_brecha.exists():
        slide.shapes.add_picture(str(img_brecha), Inches(0.5), Inches(2.55), Inches(3.7), Inches(4.0))

    # Center Column: Mapa Beta vs Drawdown (Width: 4.5 in)
    x_mid = Inches(4.42)
    w_mid = Inches(4.5)
    
    box_mid = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_mid, y_content, w_mid, h_col)
    box_mid.fill.solid()
    box_mid.fill.fore_color.rgb = RGBColor(30, 41, 59)
    box_mid.line.color.rgb = RGBColor(51, 65, 85)
    
    mtf = box_mid.text_frame
    mtf.margin_left = Inches(0.12)
    mtf.margin_top = Inches(0.1)
    mp = mtf.paragraphs[0]
    mp.text = "2. MAPA INTEGRADOR: BETA VS MAX DRAWDOWN"
    mp.font.size = Pt(11)
    mp.font.bold = True
    mp.font.color.rgb = RGBColor(167, 139, 250)
    
    img_mapa = GRAFICAS / "09_mapa_beta_drawdown.png"
    if img_mapa.exists():
        slide.shapes.add_picture(str(img_mapa), Inches(4.52), Inches(2.55), Inches(4.3), Inches(4.0))

    # Right Column: Decisiones Doña Victoria & Percentiles (Width: 3.9 in)
    x_right = Inches(9.03)
    w_right = Inches(3.9)
    
    box_right = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_right, y_content, w_right, h_col)
    box_right.fill.solid()
    box_right.fill.fore_color.rgb = RGBColor(30, 41, 59)
    box_right.line.color.rgb = RGBColor(51, 65, 85)
    
    rtf = box_right.text_frame
    rtf.word_wrap = True
    rtf.margin_left = Inches(0.12)
    rtf.margin_top = Inches(0.1)
    rp = rtf.paragraphs[0]
    rp.text = "3. DECISIÓN PARA DOÑA VICTORIA"
    rp.font.size = Pt(11)
    rp.font.bold = True
    rp.font.color.rgb = RGBColor(16, 185, 129)
    
    # Text bullets for recommendations
    bullets = [
        ("🛡️ Conservación: RECHAZAR TODO", "Evitar los 10 expedientes por extrema volatilidad e iliquidez. Asignar 100% a T-Bills (DGS3MO).", RGBColor(239, 68, 68)),
        ("📈 Crecimiento: ETF BENCHMARK", "No usar penny stocks por fragilidad operativa. Usar ETF sectoriales (XLK, XBI, XLE).", RGBColor(245, 158, 11)),
        ("⚡ Especulación: MÁX 1% - 2%", "Cupo pequeño en SOL o KOPN/TTI (alfa positivo) con stop-loss del 15% presupuestado.", RGBColor(16, 185, 129)),
    ]
    
    for b_title, b_desc, b_color in bullets:
        bp_t = rtf.add_paragraph()
        bp_t.text = b_title
        bp_t.font.size = Pt(10)
        bp_t.font.bold = True
        bp_t.font.color.rgb = b_color
        
        bp_d = rtf.add_paragraph()
        bp_d.text = b_desc
        bp_d.font.size = Pt(8.5)
        bp_d.font.color.rgb = RGBColor(226, 232, 240)
        
    img_perc = GRAFICAS / "06_matriz_percentiles.png"
    if img_perc.exists():
        slide.shapes.add_picture(str(img_perc), Inches(9.13), Inches(4.7), Inches(3.7), Inches(1.9))

    # Bottom Banner (Footer)
    footer = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.4), Inches(6.8), Inches(12.533), Inches(0.5))
    footer.fill.solid()
    footer.fill.fore_color.rgb = RGBColor(30, 41, 59)
    footer.line.color.rgb = RGBColor(51, 65, 85)
    
    ftf = footer.text_frame
    ftf.margin_left = Inches(0.2)
    ftf.margin_top = Inches(0.08)
    fp = ftf.paragraphs[0]
    fp.text = "📣 VEREDICTO FINAL: \"El 25% anual sin caídas es un espejismo. El CAPM mide la factura real del riesgo y la iliquidez.\""
    fp.font.size = Pt(10)
    fp.font.bold = True
    fp.font.color.rgb = RGBColor(248, 250, 252)

    out_pptx = OUT_DIR / "presentacion_one_pager_taller1.pptx"
    prs.save(str(out_pptx))
    print(f"PowerPoint generado con exito en: {out_pptx}")

if __name__ == "__main__":
    build_presentation()
